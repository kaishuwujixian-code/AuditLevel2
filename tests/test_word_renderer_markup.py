from docx import Document

from reporting.word_renderer import _set_paragraph_text_with_markup


def test_markup_supports_italic_links_and_red_placeholders() -> None:
    doc = Document()
    paragraph = doc.add_paragraph()

    _set_paragraph_text_with_markup(
        paragraph,
        "*Italic note with {placeholder} and https://example.com",
    )

    italic_runs = [run for run in paragraph.runs if run.text and "Italic" in run.text]
    assert italic_runs
    assert all(run.italic for run in italic_runs)

    placeholder_runs = [run for run in paragraph.runs if run.text == "{placeholder}"]
    assert placeholder_runs
    assert placeholder_runs[0].font.color.rgb is not None
    assert str(placeholder_runs[0].font.color.rgb) == "FF0000"

    xml = paragraph._p.xml
    assert "w:hyperlink" in xml
    assert "https://example.com" in xml
